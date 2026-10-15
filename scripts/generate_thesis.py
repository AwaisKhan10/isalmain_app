#!/usr/bin/env python3
"""Book-style thesis document builder for Islamian App."""

from __future__ import annotations

import lib.session
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Inches, Pt, Cm, RGBColor, Mm, Emu, Twips

ROOT = Path(__file__).resolve().parents[1]
THESIS = ROOT / "thesis"
FIG = THESIS / "figures"
DIAG = FIG / "diagrams"
SCR = FIG / "screens"
ASSETS = Path(
    "/Users/awaiskhan/.cursor/projects/Users-awaiskhan-Desktop-isalmain-app/assets"
)

NAVY = RGBColor(0x1B, 0x36, 0x5D)
MAROON = RGBColor(0x7A, 0x1F, 0x3D)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_FILL = "1B365D"

SCREEN_FILES = {
    "splash": "image-7edd5ed9-d768-4d67-afaa-5b2af6c1a998.png",
    "role": "image-60a68a12-b341-487b-b568-638f66753f94.png",
    "teacher_signin": "image-604ee26c-1234-4383-bad2-a01f9df589c3.png",
    "teacher_signup_top": "image-4cd0450a-d818-42c9-88c7-e719a247f7e7.png",
    "teacher_signup_bottom": "image-5a77e951-1cea-4238-b292-987cea9a2ccc.png",
    "student_signin": "image-409e16db-2998-4f9e-9fae-1fab98803eba.png",
    "student_signup": "image-7c005030-10cc-4d83-b5a5-7bd33b12725d.png",
    "teacher_home": "image-5ed86dde-8f73-4764-b1e9-ffc51e1bbb7f.png",
    "teacher_chat": "image-8a26f75a-51f9-4be9-9ed3-f576732a8868.png",
    "teacher_profile": "image-fb6e2cbe-451e-424f-838b-5252fcd97d5e.png",
    "teacher_profile_menu": "image-bbfe96ac-2c4d-4dde-9dfc-e8370b7c9d3a.png",
    "add_time": "image-7d0f148b-9302-4804-9077-58d418f45d4e.png",
    "add_date": "image-ae140747-ce25-4c98-aec0-c72bd2b7ab02.png",
    "add_day": "image-883e396f-5983-4708-af75-cfd43fd4f82a.png",
    "add_semester": "image-d6132934-8a8a-4cc7-87bb-17b02508fb62.png",
    "add_section": "image-2ccb8959-3036-49d4-80f6-2270f34472c2.png",
    "add_dept": "image-b3732125-3bf5-458c-b097-2c0b9322d7e8.png",
    "add_form": "image-942f7b98-d6b8-46e0-b89b-430a889a8ffb.png",
    "student_profile": "image-9a3d57b4-4b53-4dd6-a959-051042a91540.png",
    "student_chat": "image-a26d8cf2-4457-41f1-958a-d1bd86adcf62.png",
    "student_home": "image-851b21ec-d69b-472a-9194-a3b4a7d56ce4.png",
    "student_home2": "image-7f3c1457-e3c7-4fa3-bda4-9909a05073c4.png",
}


def copy_screenshots():
    SCR.mkdir(parents=True, exist_ok=True)
    for key, name in SCREEN_FILES.items():
        src = ASSETS / name
        dest = SCR / f"{key}.png"
        if src.exists():
            shutil.copy2(src, dest)
        else:
            print(f"WARNING missing screenshot: {src}")


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=BLACK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "1B365D")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def set_keep_together(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    kwn = OxmlElement("w:keepNext")
    pPr.append(kwn)


def add_page_number(paragraph, fmt="decimal"):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=11)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def set_pg_num(section, start=None, fmt="decimal"):
    sectPr = section._sectPr
    for child in list(sectPr):
        if child.tag == qn("w:pgNumType"):
            sectPr.remove(child)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    sectPr.append(pg)


def configure_section(section, left=1.5, right=1.0, top=1.15, bottom=1.15):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)


def add_header_footer(section, show_header=True, number_fmt="decimal", start=None):
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if show_header:
        r = hp.add_run("Islamian App  ·  Project Thesis  ·  Islamia College Peshawar")
        set_run_font(r, size=9, italic=True, color=NAVY)
        pPr = hp._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "1B365D")
        pBdr.append(bottom)
        pPr.append(pBdr)
    else:
        hp.text = ""

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = fp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "6")
    top.set(qn("w:space"), "4")
    top.set(qn("w:color"), "1B365D")
    pBdr.append(top)
    pPr.append(pBdr)
    add_page_number(fp)
    set_pg_num(section, start=start, fmt=number_fmt)


class Thesis:
    def __init__(self):
        self.doc = Document()
        self._setup_styles()
        configure_section(self.doc.sections[0])
        add_header_footer(self.doc.sections[0], show_header=False, number_fmt="upperRoman", start=1)
        self.doc.sections[0].different_first_page_header_footer = True

    def _setup_styles(self):
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)
        normal.font.color.rgb = BLACK
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_after = Pt(8)
        pf.space_before = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        for style_name, size, bold, space_before, space_after, align in [
            ("Heading 1", 16, True, 0, 14, WD_ALIGN_PARAGRAPH.CENTER),
            ("Heading 2", 13, True, 16, 8, WD_ALIGN_PARAGRAPH.LEFT),
            ("Heading 3", 12, True, 12, 6, WD_ALIGN_PARAGRAPH.LEFT),
        ]:
            st = styles[style_name]
            st.font.name = "Times New Roman"
            st.font.size = Pt(size)
            st.font.bold = bold
            st.font.color.rgb = NAVY
            st.font.italic = False
            st.paragraph_format.space_before = Pt(space_before)
            st.paragraph_format.space_after = Pt(space_after)
            st.paragraph_format.line_spacing = 1.15
            st.paragraph_format.alignment = align
            st.paragraph_format.keep_with_next = True

    def page_break(self):
        self.doc.add_page_break()

    def new_section(self, *, header=True, fmt="decimal", start=None):
        self.doc.add_section()
        sec = self.doc.sections[-1]
        configure_section(sec)
        add_header_footer(sec, show_header=header, number_fmt=fmt, start=start)
        return sec

    def heading1(self, chapter_no, title):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(36)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"CHAPTER {chapter_no}")
        set_run_font(r, size=14, bold=True, color=NAVY)
        p2 = self.doc.add_paragraph(style="Heading 1")
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(title.upper())
        set_run_font(r2, size=16, bold=True, color=NAVY)
        # underline bar
        p3 = self.doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(18)
        r3 = p3.add_run("—" * 18)
        set_run_font(r3, size=10, color=MAROON)

    def heading2(self, text):
        p = self.doc.add_paragraph(style="Heading 2")
        r = p.add_run(text)
        set_run_font(r, size=13, bold=True, color=NAVY)
        return p

    def heading3(self, text):
        p = self.doc.add_paragraph(style="Heading 3")
        r = p.add_run(text)
        set_run_font(r, size=12, bold=True, italic=True, color=MAROON)
        return p

    def para(self, text, *, first_indent=True, space_after=8, bold=False, center=False, size=12):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.space_after = Pt(space_after)
        if first_indent and not center:
            p.paragraph_format.first_line_indent = Cm(1.25)
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold)
        return p

    def bullet(self, text):
        p = self.doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.left_indent = Cm(1.5)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        set_run_font(r, size=12)
        return p

    def numbered(self, n, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.left_indent = Cm(1.5)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(f"{n}.  {text}")
        set_run_font(r, size=12)
        return p

    def caption(self, kind, number, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(16)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.first_line_indent = Cm(0)
        r1 = p.add_run(f"{kind} {number}:  ")
        set_run_font(r1, size=11, bold=True, color=NAVY)
        r2 = p.add_run(text)
        set_run_font(r2, size=11, italic=True)
        return p

    def figure(self, path: Path, number: str, caption: str, width=5.6):
        if not path.exists():
            self.para(f"[Figure {number} could not be loaded: {path.name}]", first_indent=False, center=True)
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run()
        run.add_picture(str(path), width=Inches(width))
        self.caption("Figure", number, caption)

    def phone(self, key: str, number: str, caption: str, width=2.35):
        self.figure(SCR / f"{key}.png", number, caption, width=width)

    def phones_row(self, keys, number, caption, width=2.15):
        table = self.doc.add_table(rows=1, cols=len(keys))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        row = table.rows[0]
        for i, key in enumerate(keys):
            cell = row.cells[i]
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            path = SCR / f"{key}.png"
            if path.exists():
                run = cell.paragraphs[0].add_run()
                run.add_picture(str(path), width=Inches(width))
        self.caption("Figure", number, caption)

    def table(self, number, caption, headers, rows, col_widths=None):
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(10)
        cap.paragraph_format.space_after = Pt(6)
        cap.paragraph_format.first_line_indent = Cm(0)
        r1 = cap.add_run(f"Table {number}:  ")
        set_run_font(r1, size=11, bold=True, color=NAVY)
        r2 = cap.add_run(caption)
        set_run_font(r2, size=11, italic=True)

        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        hdr = table.rows[0]
        prevent_row_split(hdr)
        for i, h in enumerate(headers):
            cell = hdr.cells[i]
            shade_cell(cell, HEADER_FILL)
            set_cell_border(cell)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            cell.paragraphs[0].paragraph_format.space_before = Pt(2)
            run = cell.paragraphs[0].add_run(h)
            set_run_font(run, size=10, bold=True, color=WHITE)
        for r_i, row_data in enumerate(rows):
            row = table.rows[r_i + 1]
            prevent_row_split(row)
            bg = "F4F1EA" if r_i % 2 == 0 else "FFFFFF"
            for c_i, val in enumerate(row_data):
                cell = row.cells[c_i]
                shade_cell(cell, bg)
                set_cell_border(cell)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
                cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                cell.paragraphs[0].paragraph_format.space_before = Pt(2)
                run = cell.paragraphs[0].add_run(str(val))
                set_run_font(run, size=10)
        if col_widths:
            for row in table.rows:
                for i, w in enumerate(col_widths):
                    row.cells[i].width = Inches(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(8)

    def front_title_block(self):
        d = self.doc
        for _ in range(2):
            d.add_paragraph()
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("ISLAMIA COLLEGE PESHAWAR")
        set_run_font(r, size=16, bold=True, color=NAVY)

        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Department of Computer Science")
        set_run_font(r, size=13, italic=True, color=MAROON)

        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(28)
        r = p.add_run("ISLAMIAN APP")
        set_run_font(r, size=22, bold=True, color=NAVY)

        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(
            "A Cross-Platform Mobile Application for Class Scheduling\n"
            "and Teacher–Student Communication"
        )
        set_run_font(r, size=14, italic=True)

        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(28)
        r = p.add_run("A Project Thesis submitted in partial fulfilment of the requirements\n"
                      "for the degree of Bachelor of Science in Software Engineering")
        set_run_font(r, size=12)

        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(36)
        r = p.add_run("Submitted by")
        set_run_font(r, size=12, italic=True, color=MAROON)

        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("AWAIS KHAN")
        set_run_font(r, size=16, bold=True, color=NAVY)
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Roll No.  201209")
        set_run_font(r, size=12)

        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(24)
        r = p.add_run("Supervised by")
        set_run_font(r, size=12, italic=True, color=MAROON)
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Sir Faisal Saeed")
        set_run_font(r, size=14, bold=True)

        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(48)
        r = p.add_run("Session  2020 – 2024")
        set_run_font(r, size=12, bold=True, color=NAVY)

    def inner_title(self):
        self.page_break()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(48)
        r = p.add_run("ISLAMIAN APP")
        set_run_font(r, size=20, bold=True, color=NAVY)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(
            "A Cross-Platform Mobile Application for Class Scheduling\n"
            "and Teacher–Student Communication"
        )
        set_run_font(r, size=13, italic=True)
        self.para(
            "This thesis presents the complete design and working of Islamian App, "
            "developed as a final-year project in the Department of Computer Science, "
            "Islamia College Peshawar. The work was carried out by Awais Khan under the "
            "supervision of Sir Faisal Saeed.",
            first_indent=False,
        )
        self.para("Awais Khan", center=True, first_indent=False, bold=True, space_after=2)
        self.para("Roll No. 201209", center=True, first_indent=False, space_after=16)
        self.para("Supervisor: Sir Faisal Saeed", center=True, first_indent=False)

    def certificate(self):
        self.page_break()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("CERTIFICATE OF APPROVAL")
        set_run_font(r, size=16, bold=True, color=NAVY)
        self.para(
            "This is to certify that the work contained in this thesis entitled "
            "“Islamian App: A Cross-Platform Mobile Application for Class Scheduling "
            "and Teacher–Student Communication” was carried out by Awais Khan "
            "(Roll No. 201209) under my supervision. In my opinion the work is adequate, "
            "in scope and quality, for the award of the degree of Bachelor of Science "
            "in Software Engineering from Islamia College Peshawar.",
            first_indent=False,
        )
        self.para("Approved by:", first_indent=False, bold=True)
        self.para("Signature: ______________________________", first_indent=False)
        self.para("Supervisor: Sir Faisal Saeed", first_indent=False)
        self.para("Department of Computer Science, Islamia College Peshawar", first_indent=False)
        self.para("Examiner: ______________________________", first_indent=False)
        self.para("Chairman / Coordinator: Dr. Atif Khan", first_indent=False)
        self.para("Department of Computer Science, Islamia College Peshawar", first_indent=False)

    def dedication(self):
        self.page_break()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("DEDICATION")
        set_run_font(r, size=16, bold=True, color=NAVY)
        self.para(
            "This work is dedicated to my parents, whose prayers and sacrifices made "
            "this journey possible, and to my teachers, whose guidance shaped both this "
            "project and my understanding of software engineering.",
            first_indent=False, center=True,
        )

    def acknowledgement(self):
        self.page_break()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("ACKNOWLEDGEMENT")
        set_run_font(r, size=16, bold=True, color=NAVY)
        self.para(
            "First and foremost, I thank Allah Almighty for the strength and patience "
            "required to complete this project. Blessings be upon the Prophet Muhammad "
            "(peace be upon him)."
        )
        self.para(
            "I am sincerely grateful to my supervisor, Sir Faisal Saeed, for his careful "
            "reading of this work, for pointing out where the writing had to become clearer, "
            "and for insisting that a thesis should read like a book rather than a collection "
            "of technical notes. His comments on diagrams, alignment, and the explanation of "
            "each screen guided the present version of this report."
        )
        self.para(
            "I also thank the faculty of the Department of Computer Science, Islamia College "
            "Peshawar, and my family, whose support remained constant throughout the degree."
        )

    def abstract(self):
        self.page_break()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("ABSTRACT")
        set_run_font(r, size=16, bold=True, color=NAVY)
        self.para(
            "Class timetables in many colleges are still shared through notice boards, "
            "paper sheets, or informal chat groups. These methods scatter information, "
            "go out of date quickly, and give students no reliable way to speak to the "
            "teacher who owns a particular class. Islamian App was designed to remove "
            "this gap by placing timetable publication and teacher–student conversation "
            "inside one mobile application."
        )
        self.para(
            "The system has two roles. A teacher registers, signs in, and publishes class "
            "meetings by choosing department, field of study, section, subject, semester, "
            "day, date and time. A student registers with the same academic details and "
            "then sees only those classes that match his or her profile. From a class card "
            "the student may open a private chat with that teacher. The teacher, in turn, "
            "can reply from the chat tab of the application."
        )
        self.para(
            "The application was built as a cross-platform mobile client. User accounts "
            "are verified through a cloud authentication service. Profiles, schedules and "
            "messages are stored in a cloud database so that a change made by a teacher "
            "appears on the student’s screen without a separate announcement. This thesis "
            "describes the problem, the plan of work, the requirements, the design of the "
            "database and the diagrams of the system, and then walks through every screen "
            "of both the teacher flow and the student flow in the order a new user would "
            "meet them."
        )
        self.para("Keywords: class scheduling, teacher–student communication, mobile application, "
                  "academic timetable, Islamia College Peshawar.", first_indent=False, size=11)

    def toc_placeholder(self):
        self.page_break()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("TABLE OF CONTENTS")
        set_run_font(r, size=16, bold=True, color=NAVY)
        note = self.doc.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = note.add_run("(Update in Microsoft Word: right-click the list below → Update Field)")
        set_run_font(rr, size=9, italic=True, color=MAROON)
        p = self.doc.add_paragraph()
        run = p.add_run()
        fld1 = OxmlElement("w:fldChar")
        fld1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = r' TOC \o "1-3" \h \z \u '
        fld2 = OxmlElement("w:fldChar")
        fld2.set(qn("w:fldCharType"), "separate")
        t = OxmlElement("w:t")
        t.text = "Right-click here and choose Update Field to build the table of contents."
        fld3 = OxmlElement("w:fldChar")
        fld3.set(qn("w:fldCharType"), "end")
        run._r.append(fld1)
        run._r.append(instr)
        run._r.append(fld2)
        run._r.append(t)
        run._r.append(fld3)

    def lists_of_tables_figures(self):
        self.page_break()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("LIST OF TABLES")
        set_run_font(r, size=16, bold=True, color=NAVY)
        tables = [
            "2.1  Technology stack of the project",
            "2.2  Stakeholder analysis",
            "2.3  Feasibility summary",
            "3.1  Functional requirements",
            "3.2  Non-functional requirements",
            "3.3  Requirements traceability matrix",
            "4.1  Teacher entity — data dictionary",
            "4.2  Student entity — data dictionary",
            "4.3  Class schedule entity — data dictionary",
            "4.4  Conversation entity — data dictionary",
            "4.5  Message entity — data dictionary",
            "4.6  Allowed values used in drop-down lists",
            "6.1  Test cases and results",
        ]
        for t in tables:
            pp = self.doc.add_paragraph()
            pp.paragraph_format.space_after = Pt(2)
            pp.paragraph_format.first_line_indent = Cm(0)
            run = pp.add_run(t)
            set_run_font(run, size=12)

        self.page_break()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("LIST OF FIGURES")
        set_run_font(r, size=16, bold=True, color=NAVY)
        figs = [
            "4.1  System architecture",
            "4.2  Use-case diagram",
            "4.3  Context data-flow diagram (Level 0)",
            "4.4  Entity relationship diagram",
            "4.5  Activity diagram — publishing a class schedule",
            "4.6  Deployment diagram",
            "4.7  Screen navigation map",
            "5.1  Splash screen",
            "5.2  Role-selection screen",
            "5.3  Teacher sign-in screen",
            "5.4  Teacher sign-up screen",
            "5.5  Teacher home — My Class Schedules",
            "5.6  Add Class Schedule form",
            "5.7  Selecting department, section, semester and day",
            "5.8  Selecting class date and class time",
            "5.9  Teacher chat screen",
            "5.10 Teacher profile screen",
            "5.11 Student sign-in screen",
            "5.12 Student sign-up screen",
            "5.13 Student home — My Classes",
            "5.14 Student chat with a teacher",
            "5.15 Student profile screen",
        ]
        for t in figs:
            pp = self.doc.add_paragraph()
            pp.paragraph_format.space_after = Pt(2)
            pp.paragraph_format.first_line_indent = Cm(0)
            run = pp.add_run(t)
            set_run_font(run, size=12)

        self.page_break()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("LIST OF ABBREVIATIONS")
        set_run_font(r, size=16, bold=True, color=NAVY)
        abbr = [
            ("ERD", "Entity Relationship Diagram"),
            ("DFD", "Data Flow Diagram"),
            ("GUI", "Graphical User Interface"),
            ("PK", "Primary Key"),
            ("FK", "Foreign Key"),
            ("HTTPS", "Hypertext Transfer Protocol Secure"),
            ("UI", "User Interface"),
        ]
        self.table("0.1", "Abbreviations used in this thesis",
                   ["Abbreviation", "Meaning"], [[a, b] for a, b in abbr])

    def ch1(self):
        self.heading1("ONE", "Introduction")
        self.heading2("1.1  Overview")
        self.para(
            "A college timetable is useful only when the right student sees the right class "
            "at the right time, and when that student can still reach the teacher if something "
            "goes wrong. In daily campus life this simple need is often met in a clumsy way. "
            "A teacher writes a time on the board, forwards a photograph of a timetable into "
            "a messaging group, or tells a class representative to spread the word. Students "
            "who were absent, who changed section, or who belong to another group never receive "
            "the update. When they need to ask a question they must hunt for the teacher’s "
            "personal number, which many teachers are unwilling to share."
        )
        self.para(
            "Islamian App is a mobile application written for this setting. It belongs to "
            "Islamia College Peshawar in name and purpose: it gives teachers a place to publish "
            "class meetings and gives students a place to read only those meetings that belong "
            "to their department, field, section and semester. Beside the timetable it provides "
            "a private conversation so that a student can, for example, inform a teacher of "
            "illness without leaving the application."
        )
        self.para(
            "The remainder of this chapter states why the work was undertaken, what problem "
            "it attacks, what solution is offered, and where the boundaries of the project lie."
        )

        self.heading2("1.2  Motivation")
        self.para(
            "Two observations motivated the project. First, almost every student and teacher "
            "already carries a smartphone, yet the official timetable still lives on paper or "
            "in unstructured chats. Second, conversation about a class is most useful when it "
            "is tied to that class, not scattered across personal WhatsApp threads. A dedicated "
            "campus application can keep both the timetable and the conversation in one trusted "
            "place, under the identity of the college rather than a private number."
        )

        self.heading2("1.3  Problem Statement")
        self.para(
            "There is no single, role-based digital system in the target setting that allows a "
            "teacher to publish a structured class meeting and allows a student to see only the "
            "meetings that match his or her academic profile, while still offering a direct, "
            "in-application channel of communication between the two."
        )
        self.para(
            "Paper notices cannot be updated instantly. General-purpose chat applications do "
            "not know a student’s section or semester, so they cannot hide irrelevant classes. "
            "As a result students either miss changes or receive noise, and teachers cannot "
            "be sure that the correct audience has been reached."
        )

        self.heading2("1.4  Proposed Solution")
        self.para(
            "The proposed solution is a dual-role mobile application. After a short splash "
            "screen the user is asked whether he or she is a teacher or a student. From that "
            "choice the application opens the matching sign-in page. A new user may register. "
            "A teacher who has signed in lands on a home screen titled My Class Schedules and "
            "may add, edit or delete meetings. A student who has signed in lands on a home "
            "screen titled My Classes and sees a filtered list. Both roles share a chat tab "
            "and a profile tab. The next screen after each tap is described in Chapter 5, "
            "together with the actual photographs of the running application."
        )

        self.heading2("1.5  Objectives")
        self.numbered(1, "To study the difficulties of sharing class times by informal means and to state a clear problem.")
        self.numbered(2, "To design a two-role application in which teachers publish schedules and students consume a personal timetable.")
        self.numbered(3, "To provide a private chat between a student and the teacher of a class.")
        self.numbered(4, "To store users, schedules and messages in a cloud database so that updates appear in real time.")
        self.numbered(5, "To document the system in the manner of a book: with diagrams, data tables, and a screen-by-screen explanation of the working product.")

        self.heading2("1.6  Scope")
        self.para("The project includes the following:")
        self.bullet("Mobile application for Android (and other phones that can run the same client).")
        self.bullet("Registration and sign-in for teachers and for students.")
        self.bullet("Creation, editing and deletion of class schedules by teachers.")
        self.bullet("Filtered timetable for students, based on department, field, section and semester.")
        self.bullet("One-to-one chat, profile viewing and profile editing, and sign-out.")
        self.para(
            "The project does not include a web administration panel, attendance marking, "
            "fee payment, video lectures, classroom maps, or automatic generation of a clash-free "
            "timetable for the whole college. Those items remain recommendations for later work."
        )

        self.heading2("1.7  Limitations")
        self.bullet("The application requires an internet connection; there is no full offline mode.")
        self.bullet("Push notifications on a locked phone are not yet delivered; the user sees new messages after opening the chat.")
        self.bullet("There is no college-wide administrator who can upload an entire semester timetable in one file.")
        self.bullet("The present design assumes a single campus and a known list of departments.")

        self.heading2("1.8  Organisation of the Thesis")
        self.para(
            "Chapter 2 records the plan of the project, the people involved, and a short "
            "feasibility study. Chapter 3 states the requirements in ordinary language. "
            "Chapter 4 presents the architecture, the use-case view, the data-flow view, "
            "the entity-relationship diagram and the other design drawings, together with "
            "data tables that match the live system. Chapter 5 is the heart of the report: "
            "it walks from the first splash screen to the last profile screen and explains "
            "what happens when each button is pressed. Chapter 6 summarises testing. "
            "Chapter 7 concludes and points to future work."
        )

        self.heading2("1.9  Summary")
        self.para(
            "Islamian App is a campus mobile application that joins two everyday academic "
            "tasks — publishing a class time and speaking to the teacher of that class — "
            "in one place. The following chapters treat the work as a finished product to "
            "be read, not as a pile of source files to be decoded."
        )

    def ch2(self):
        self.page_break()
        self.heading1("TWO", "Project Planning")
        self.heading2("2.1  Project Identity")
        self.para(
            "The product is named Islamian App. It is a class-scheduling and messaging "
            "application for teachers and students. Development followed short cycles: "
            "first the accounts, then the timetable, then the chat, and finally the profile "
            "screens. Each cycle was tried on a real phone before the next cycle began."
        )

        self.heading3("2.1.1  Main features")
        self.bullet("A splash screen that shows the college emblem while the application prepares itself.")
        self.bullet("A role-selection screen that asks whether the user is a teacher or a student.")
        self.bullet("Separate sign-in and sign-up pages for each role.")
        self.bullet("A teacher home on which class cards can be added, edited or removed.")
        self.bullet("A student home that lists only matching classes, each with a shortcut to chat.")
        self.bullet("A conversation screen with timestamps and an online / offline label.")
        self.bullet("Profile pages that display the stored academic and personal details.")

        self.heading3("2.1.2  Tools and technologies")
        self.para(
            "Table 2.1 lists only those tools that the reader needs in order to understand "
            "the finished system. Brand names of programming libraries are omitted from the "
            "narrative; they belong in an appendix, not in the argument of the thesis."
        )
        self.table("2.1", "Technology stack of the project",
                   ["Item", "Role in this project"],
                   [
                       ["Flutter (Dart)", "Builds the screens that the teacher and the student see."],
                       ["Firebase Authentication", "Creates and checks email-and-password accounts."],
                       ["Cloud Firestore", "Holds teacher records, student records, class schedules and chat data."],
                       ["Cloud Storage", "Holds profile photographs."],
                       ["On-device preference store", "Remembers that the opening tutorial has already been seen."],
                   ])

        self.heading3("2.1.3  Method of work")
        self.para(
            "An iterative method was used rather than a single long waterfall. Authentication "
            "was completed first so that every later screen had a signed-in user. Scheduling "
            "came next because it is the reason the application exists. Chat was added once "
            "a teacher identity could be attached to a class card. Profiles and polishing "
            "occupied the last cycle. This order matches the way a user actually travels "
            "through the product."
        )

        self.heading3("2.1.4  Risks")
        self.bullet("If teachers do not keep schedules up to date, students will stop opening the home screen.")
        self.bullet("If the cloud service is misconfigured, personal data could be exposed.")
        self.bullet("If the lists of departments and fields become outdated, new programmes cannot be selected.")
        self.bullet("If the internet is unavailable during a class change, the student will not see the new time.")

        self.heading2("2.2  Stakeholders")
        self.table("2.2", "People who have an interest in the system",
                   ["Stakeholder", "Interest", "How the system serves them"],
                   [
                       ["Teacher", "To announce classes and answer students.", "Home, add-schedule form, chat, profile."],
                       ["Student", "To know when a class meets and to message the teacher.", "Filtered home, chat shortcut, profile."],
                       ["Supervisor / examiner", "To judge completeness and clarity of the work.", "This thesis, diagrams and screen record."],
                       ["College administration", "To keep academic communication orderly.", "A named campus application instead of private chats."],
                   ])

        self.heading2("2.3  Feasibility")
        self.para(
            "The work is technically feasible because ready-made cloud services already "
            "provide accounts, a database and file storage, while a single mobile toolkit "
            "draws both the teacher and the student screens. It is economically feasible "
            "because the project used student licences and a free cloud tier sufficient "
            "for demonstration. It is operationally feasible because the steps a user must "
            "take are the same steps already known from everyday applications: choose a "
            "role, sign in, tap a card, send a message. The calendar of a final-year "
            "project is tight, yet the scope was held to timetable and chat so that the "
            "product could be finished rather than merely planned."
        )
        self.table("2.3", "Summary of feasibility",
                   ["Question", "Finding"],
                   [
                       ["Can it be built with the available skills and tools?", "Yes. The student developer already used the chosen mobile toolkit, and the cloud services are documented."],
                       ["Can it be afforded?", "Yes. Development cost is the student’s time; running cost at demonstration scale is within the free cloud allowance."],
                       ["Will people be able to use it?", "Yes. Each screen asks for ordinary academic facts (department, section, time) in drop-down lists."],
                       ["Can it be finished in one academic session?", "Yes, provided the scope stays at scheduling plus chat, which it did."],
                   ])

        self.heading2("2.4  Summary")
        self.para(
            "Planning fixed the name, the two roles, the four main features, the tools, "
            "the risks and the people who must be satisfied. The next chapter turns those "
            "plans into numbered requirements."
        )

    def ch3(self):
        self.page_break()
        self.heading1("THREE", "Requirement Engineering")
        self.heading2("3.1  Business requirement")
        self.para(
            "The college needs a single mobile place in which a teacher can publish a class "
            "meeting and a student of that meeting can read it and, if necessary, write to "
            "the teacher."
        )
        self.heading2("3.2  User requirements")
        self.para("A teacher must be able to:")
        self.bullet("Create an account with name, email, phone, department, qualification, subjects and gender.")
        self.bullet("Sign in and reach a list of the classes he or she has published.")
        self.bullet("Add a class by filling department, field, section, subject, semester, day, date and time.")
        self.bullet("Correct or delete a class, and chat with a student.")
        self.para("A student must be able to:")
        self.bullet("Create an account with name, email, department, field, section and semester.")
        self.bullet("Sign in and see only the classes that match those academic facts.")
        self.bullet("Open a chat with the teacher named on a class card.")
        self.bullet("View and update a personal profile, and sign out.")

        self.heading2("3.3  Functional requirements")
        self.table("3.1", "Functional requirements",
                   ["ID", "Requirement", "Who uses it"],
                   [
                       ["FR-01", "Show a splash screen, then a role-selection screen.", "Both"],
                       ["FR-02", "Register a teacher and store the professional profile.", "Teacher"],
                       ["FR-03", "Register a student and store the academic profile.", "Student"],
                       ["FR-04", "Sign in with email and password, or with a Google account.", "Both"],
                       ["FR-05", "Keep the user signed in until he or she chooses Log Out.", "Both"],
                       ["FR-06", "Let a teacher add, edit and delete a class schedule.", "Teacher"],
                       ["FR-07", "Refuse to save a schedule if date or time is missing.", "Teacher"],
                       ["FR-08", "Show a student only schedules that match department, field, section and semester.", "Student"],
                       ["FR-09", "Open a private chat from a class card or from the Chat tab.", "Both"],
                       ["FR-10", "Show online or offline status of the other person in a chat.", "Both"],
                       ["FR-11", "Display and edit the user’s profile.", "Both"],
                       ["FR-12", "Provide Home, Chat and Profile tabs after sign-in.", "Both"],
                   ])

        self.heading2("3.4  Non-functional requirements")
        self.table("3.2", "Non-functional requirements",
                   ["ID", "Quality", "Expectation"],
                   [
                       ["NFR-01", "Usability", "A new user should complete sign-up without training, using labelled fields and drop-down lists."],
                       ["NFR-02", "Performance", "A saved schedule should appear on the other role’s list within a few seconds on a normal connection."],
                       ["NFR-03", "Security", "Passwords are checked by the authentication service; only a signed-in user reaches Home, Chat and Profile."],
                       ["NFR-04", "Availability", "The application depends on the cloud service being reachable."],
                       ["NFR-05", "Maintainability", "Teacher screens and student screens are separate, so a change to one role need not break the other."],
                   ])

        self.heading2("3.5  Use-case narratives")
        self.heading3("3.5.1  Teacher adds a class")
        self.para(
            "The teacher is already on My Class Schedules. He taps the round plus button. "
            "A form titled Add Class Schedule slides up. He chooses Computer Science as "
            "department, Software Engineering as field, section B, types a subject name, "
            "chooses a semester, a weekday, a calendar date and a clock time, then taps "
            "Submit Schedule. If any required value is missing, a red sentence appears "
            "under that field and the form does not close. If everything is present, the "
            "form closes and a new coloured card is added to the list."
        )
        self.heading3("3.5.2  Student opens a chat")
        self.para(
            "The student is on My Classes and sees a card for a subject taught by a named "
            "teacher. He taps the small speech-bubble on that card. The chat screen opens, "
            "showing the teacher’s name and whether the teacher is online. He types a "
            "message and taps the send button. The bubble appears on the right. When the "
            "teacher later opens the same conversation, the same words appear on the left."
        )

        self.heading2("3.6  Traceability")
        self.table("3.3", "How requirements are shown later in this thesis",
                   ["Requirement", "Design", "Working screen"],
                   [
                       ["FR-01", "Figure 4.7", "Figures 5.1 and 5.2"],
                       ["FR-02, FR-03, FR-04", "Use cases Register / Sign in", "Figures 5.3, 5.4, 5.11, 5.12"],
                       ["FR-06, FR-07", "ERD Class Schedule; activity diagram", "Figures 5.5 to 5.8"],
                       ["FR-08", "Dashed relationship on the ERD", "Figure 5.13"],
                       ["FR-09, FR-10", "Conversation and Message entities", "Figures 5.9 and 5.14"],
                       ["FR-11, FR-12", "Navigation map", "Figures 5.10 and 5.15"],
                   ])

        self.heading2("3.7  Summary")
        self.para(
            "The requirements are modest in number and can each be pointed to a screen. "
            "Chapter 4 now draws the system that satisfies them; Chapter 5 shows the screens "
            "themselves."
        )

    def ch4(self):
        self.page_break()
        self.heading1("FOUR", "System Modelling and Design")
        self.heading2("4.1  Architecture")
        self.para(
            "The application is arranged in three layers, as drawn in Figure 4.1. At the "
            "top sit the screens: one family for the teacher and one family for the student. "
            "In the middle sit four jobs that the screens share — checking who the user is, "
            "saving and reading class times, sending messages, and keeping the profile. "
            "At the bottom sit the cloud services: an authentication service for accounts, "
            "a database for records, a file store for photographs, and a small on-device "
            "memory that remembers whether the user has already seen the first-launch pages."
        )
        self.para(
            "A teacher never sees a student form, and a student never sees the add-schedule "
            "button. The split is made as soon as the role-selection screen is answered, and "
            "it is confirmed again when the account is created, because a teacher record and "
            "a student record are different things in the database."
        )
        self.figure(DIAG / "fig_4_1_architecture.png", "4.1",
                    "Three-layer architecture of Islamian App.")

        self.heading2("4.2  Use-case view")
        self.para(
            "Figure 4.2 shows the two actors and the work they can ask of the system. "
            "Both actors register, sign in, manage a profile, exchange messages and sign out. "
            "Only the teacher creates, updates and deletes a class schedule. Only the student "
            "reads a personal timetable and starts a chat from a class card. The ovals inside "
            "the rectangle are the system; the stick figures outside it are people."
        )
        self.figure(DIAG / "fig_4_2_usecase.png", "4.2",
                    "Use-case diagram. Shared actions sit in the middle; role-specific actions sit lower down.")

        self.heading2("4.3  Data flow")
        self.para(
            "Figure 4.3 is a context diagram. The teacher sends schedule details and messages "
            "into the system and receives confirmations and student replies. The student sends "
            "a profile and chat text and receives a personal timetable. Inside the system the "
            "flows settle into three stores: user records, class schedules, and conversations. "
            "No store is written by a person directly; every arrow passes through the application."
        )
        self.figure(DIAG / "fig_4_3_dfd.png", "4.3",
                    "Level-0 data flow (context) diagram.")

        self.heading2("4.4  Entity relationship diagram")
        self.para(
            "The live database contains five kinds of record. They are drawn in Figure 4.4 "
            "and tabulated in this section so that the picture and the documentation describe "
            "the same product."
        )
        self.para(
            "A Teacher creates many Class Schedule rows. Each schedule stores the teacher’s "
            "identifier so that the home list of that teacher can be rebuilt, and also stores "
            "the teacher’s name so that a student card can show “Teacher: …” without a second "
            "look-up. A Student does not hold a list of schedule identifiers. Instead, when "
            "the student opens My Classes, the application asks the database for those schedules "
            "whose department, field, section and semester are the same as the student’s own. "
            "That rule is drawn as a dashed line, because it is a match of values rather than "
            "a stored foreign key."
        )
        self.para(
            "Chat is modelled in two steps. A Conversation remembers the two people, the last "
            "line of text, the time of that line, and how many unread lines each person still "
            "has. A Message belongs to one conversation and records who sent it, who should "
            "read it, the words, and the time. One conversation therefore contains many messages."
        )
        self.figure(DIAG / "fig_4_4_erd.png", "4.4",
                    "Entity relationship diagram. Primary keys are marked PK; foreign keys are marked FK. "
                    "The dashed line is a logical match, not a stored link.")

        self.heading3("4.4.1  Data dictionary")
        self.para(
            "The following tables are the documentation of the five entities. Attribute names "
            "are written in plain English. They correspond one-for-one with the fields stored "
            "for the running application."
        )
        self.table("4.1", "Teacher entity",
                   ["Attribute", "Type", "Description"],
                   [
                       ["teacher_id", "Text (key)", "Unique identity of the teacher, issued at registration."],
                       ["full_name", "Text", "Name shown on class cards and in chat, e.g. Zubair."],
                       ["email", "Text", "Sign-in address, e.g. zubair@gmail.com."],
                       ["phone_no", "Text", "Contact number shown on the profile, e.g. 03039639571."],
                       ["department", "Text", "Chosen from the college list, e.g. Computer Science."],
                       ["qualification", "Text", "Highest degree entered by the teacher, e.g. PhD."],
                       ["subjects", "Text", "Subjects the teacher offers, e.g. Computer Science."],
                       ["gender", "Text", "Selected from Male / Female / Other."],
                       ["profile_image", "Text (link)", "Photograph stored in the file service, if any."],
                       ["is_online", "Yes / No", "Whether the teacher currently has the application open."],
                   ])
        self.table("4.2", "Student entity",
                   ["Attribute", "Type", "Description"],
                   [
                       ["student_id", "Text (key)", "Unique identity of the student, issued at registration."],
                       ["full_name", "Text", "Name shown in chat, e.g. Awais Khan."],
                       ["email", "Text", "Sign-in address, e.g. awaiskhan@gmail.com."],
                       ["department", "Text", "Must match a teacher’s department for classes to appear."],
                       ["field", "Text", "Programme inside the department, e.g. Software Engineering."],
                       ["section", "Text", "A, B, C or D."],
                       ["semester", "Text", "1 to 8."],
                       ["profile_image", "Text (link)", "Optional photograph."],
                       ["is_online", "Yes / No", "Whether the student currently has the application open."],
                   ])
        self.table("4.3", "Class schedule entity",
                   ["Attribute", "Type", "Description"],
                   [
                       ["schedule_id", "Text (key)", "Unique identity of one published meeting."],
                       ["teacher_id", "Text (FK)", "The teacher who created the meeting."],
                       ["teacher_name", "Text", "Copy of the teacher’s name for the student card."],
                       ["department", "Text", "Audience filter, e.g. Computer Science."],
                       ["field", "Text", "Audience filter, e.g. Software Engineering."],
                       ["class_section", "Text", "Audience filter: A, B, C or D."],
                       ["subject", "Text", "Title typed by the teacher, e.g. Calculus, Islamiyat."],
                       ["semester", "Text", "Audience filter: 1 to 8."],
                       ["day", "Text", "Weekday chosen from Monday to Sunday."],
                       ["class_date", "Date", "Calendar day of the meeting."],
                       ["class_time", "Time", "Clock time of the meeting, e.g. 3:19 PM."],
                   ])
        self.table("4.4", "Conversation entity",
                   ["Attribute", "Type", "Description"],
                   [
                       ["conversation_id", "Text (key)", "Built from the two user identities so that the same pair always reopens the same thread."],
                       ["participant_ids", "List", "The teacher identity and the student identity."],
                       ["participant_names", "Text", "Names cached so the chat list can be drawn quickly."],
                       ["last_message", "Text", "Most recent line, shown as a preview."],
                       ["last_time", "Date-time", "When that line was sent."],
                       ["unread_counts", "Map", "How many new lines each participant has not yet opened."],
                   ])
        self.table("4.5", "Message entity",
                   ["Attribute", "Type", "Description"],
                   [
                       ["message_id", "Text (key)", "Unique identity of one bubble."],
                       ["conversation_id", "Text (FK)", "Thread to which the bubble belongs."],
                       ["sender_id", "Text", "Who typed the line."],
                       ["receiver_id", "Text", "Who should read it."],
                       ["content", "Text", "The words, including ordinary greeting or a leave request."],
                       ["sent_time", "Date-time", "Shown under the bubble, e.g. 08:05 AM."],
                   ])
        self.table("4.6", "Controlled lists used by the forms",
                   ["List", "Allowed values"],
                   [
                       ["Department", "Computer Science; Business Administration; Electrical Engineering; Mathematics; Physics; English."],
                       ["Field (depends on department)", "For Computer Science: Software Engineering, Computer Engineering, Artificial Intelligence, Information Technology, Data Science, Cyber Security, Management Information Systems. Other departments have their own programme lists."],
                       ["Section", "A, B, C, D."],
                       ["Semester", "1, 2, 3, 4, 5, 6, 7, 8."],
                       ["Day", "Monday through Sunday."],
                   ])
        self.para(
            "These lists are the reason the forms in Chapter 5 use drop-down menus rather "
            "than free typing for department, field, section, semester and day. Free typing "
            "would break the matching rule that connects a student to a schedule."
        )

        self.heading2("4.5  Activity: publishing a class")
        self.para(
            "Figure 4.5 follows a teacher from the home list to a saved card. The diamond "
            "in the middle is the same check that produces the red sentences “Date is required” "
            "and “Time is required” on the form. A failed check returns the teacher to the "
            "form; a successful check writes the row and shows the new card."
        )
        self.figure(DIAG / "fig_4_5_activity.png", "4.5",
                    "Activity diagram for adding a class schedule.", width=5.2)

        self.heading2("4.6  Deployment")
        self.para(
            "Figure 4.6 places the finished product on real machines. The client is the "
            "student’s or teacher’s phone. The cloud node is the authentication service, "
            "the database and the file store. Traffic between them travels over HTTPS."
        )
        self.figure(DIAG / "fig_4_6_deployment.png", "4.6",
                    "Deployment of the mobile client and the cloud services.")

        self.heading2("4.7  Navigation of screens")
        self.para(
            "Figure 4.7 is a map of every screen that Chapter 5 will photograph. It is the "
            "reading order of the next chapter. The left branch is the teacher; the right "
            "branch is the student. After sign-in, three tabs — Home, Chat and Profile — "
            "are always available at the bottom of the screen."
        )
        self.figure(DIAG / "fig_4_7_navigation.png", "4.7",
                    "How a user travels from the splash screen to the working tabs.")

        self.heading2("4.8  Summary")
        self.para(
            "The design is small on purpose. Five entities, two actors, and a handful of "
            "screens are enough to tell the story of a class that is published and a message "
            "that follows it. The next chapter opens the application and walks through those "
            "screens in order, with photographs taken from the running product."
        )

    def ch5(self):
        self.page_break()
        self.heading1("FIVE", "Implementation and User Interface")
        self.heading2("5.1  How this chapter is to be read")
        self.para(
            "This chapter does not explain source files. It explains the product the way a "
            "book explains a tool: screen by screen, in the order a new person meets them. "
            "For each screen three things are stated. First, what the screen is for. Second, "
            "what the user sees. Third, what happens when a control is used, and which screen "
            "opens next. Photographs of the live application are placed immediately under the "
            "prose they illustrate."
        )
        self.para(
            "The walk begins with the two screens that both roles share. It then follows the "
            "teacher from sign-in to profile. It then follows the student along the parallel "
            "path. A reader who only wishes to understand the student experience may start at "
            "Section 5.4."
        )

        self.heading2("5.2  Screens common to both roles")
        self.heading3("5.2.1  Splash screen")
        self.para(
            "When the application is tapped on the phone, the first picture that appears is "
            "the emblem of Islamia College Peshawar on a white field. Nothing is asked of the "
            "user. The pause is short. Behind the emblem the application is doing quiet work: "
            "it checks whether someone is already signed in, and whether this phone has seen "
            "the opening pages before."
        )
        self.para(
            "If a teacher session is still alive, the next screen is My Class Schedules. If a "
            "student session is still alive, the next screen is My Classes. If nobody is signed "
            "in, the next screen is the role-selection page shown in Figure 5.2. The emblem "
            "therefore has a double duty: it names the college, and it covers the moment in "
            "which the application decides where to go."
        )
        self.phone("splash", "5.1",
                   "Splash screen. The college emblem is shown while the application decides the next page.")

        self.heading3("5.2.2  Role-selection screen")
        self.para(
            "The second screen asks a single question, written in capital letters: “YOU ARE "
            "TEACHER OR STUDENT”. Above the question stands a line drawing of the college "
            "building and the words Islamia College University Peshawar. Below the question "
            "stand two large buttons. The upper button is labelled Teacher and carries a small "
            "picture of an instructor at a board. The lower button is labelled Student and "
            "carries a graduation cap."
        )
        self.para(
            "This choice is the fork of the whole product. Tapping Teacher opens the teacher "
            "sign-in page (Figure 5.3). Tapping Student opens the student sign-in page "
            "(Figure 5.11). The two roads do not cross again except in a chat, where a teacher "
            "and a student meet as two ends of one conversation. There is no third role, and "
            "there is no way to open both homes at once on the same account."
        )
        self.phone("role", "5.2",
                   "Role-selection screen. The Teacher button opens teacher sign-in; the Student button opens student sign-in.")

        self.heading2("5.3  Teacher flow")
        self.heading3("5.3.1  Teacher sign-in")
        self.para(
            "The sign-in page is titled Sign In. A back arrow at the top left returns the user "
            "to the role-selection screen, in case Teacher was tapped by mistake. Two fields "
            "ask for an email address and a password. An eye icon on the password field shows "
            "or hides the letters. The wide gradient button labelled Sign In checks the "
            "account. If the email and password belong to a registered teacher, the next "
            "screen is My Class Schedules (Figure 5.5)."
        )
        self.para(
            "A second button, Sign In Using Google, offers the same destination by means of "
            "a Google account. Under the buttons a short line reads “New Here? Sign Up”. "
            "Tapping the blue words Sign Up opens the teacher registration form in Figure 5.4. "
            "A person who is already registered never needs that form; a person who is not "
            "registered cannot pass the Sign In button."
        )
        self.phone("teacher_signin", "5.3",
                   "Teacher sign-in. Success opens My Class Schedules; Sign Up opens the registration form; the back arrow returns to role selection.")

        self.heading3("5.3.2  Teacher sign-up")
        self.para(
            "The registration page is titled Teacher Sign Up. Its instruction is plain: "
            "“Please fill the form fields with the following information.” The form collects "
            "eight facts that later appear on the teacher profile and on student class cards."
        )
        self.bullet("Full Name — the name students will see, for example Zubair.")
        self.bullet("Email address — the address that will be used to sign in afterwards.")
        self.bullet("Phone Number — stored on the profile, for example 03039639571.")
        self.bullet("Select Department — a drop-down list of the six college departments.")
        self.bullet("Qualification — typed freely, for example Masters or PhD.")
        self.bullet("Subjects — typed freely, for example Maths, CS.")
        self.bullet("Select Gender — a drop-down list.")
        self.bullet("Password — hidden by default, with an eye icon to reveal it.")
        self.para(
            "When Register is tapped and the fields are acceptable, the account is created "
            "and the teacher is taken to My Class Schedules. If the teacher already has an "
            "account, the line “Already have an account? Sign In” returns to Figure 5.3. "
            "The back arrow likewise returns to sign-in. Nothing on this page creates a class; "
            "it only creates the person who will later create classes."
        )
        self.phones_row(["teacher_signup_top", "teacher_signup_bottom"], "5.4",
                        "Teacher sign-up. The upper part names the page; the lower part holds the remaining fields, the Register button, and the link back to Sign In.")

        self.heading3("5.3.3  Teacher home — My Class Schedules")
        self.para(
            "After a successful sign-in the teacher arrives at the home tab. The title in "
            "the centre reads My Class Schedules. A circular arrow at the top right refreshes "
            "the list. Each published class occupies a dark card that fades from maroon to "
            "navy. On the cards in Figure 5.5 the subjects are Calculus and Islamiyat. Each "
            "card states the department (Computer Science), the field (Software Engineering), "
            "the section (B), the semester (8), the weekday, the calendar date and the clock "
            "time. Two small icons sit at the top right of a card: a pencil to edit, and a "
            "bin to delete."
        )
        self.para(
            "Tapping the pencil reopens the same form that was used to create the class, "
            "already filled with the old values, so that a time can be corrected without "
            "typing everything again. Tapping the bin asks the teacher to confirm, then "
            "removes the card; students who matched that class will no longer see it after "
            "their list refreshes. Tapping the round plus button at the lower right opens "
            "the Add Class Schedule form described in the next section."
        )
        self.para(
            "The bar at the bottom of the screen has three tabs. Home is the present page "
            "and is drawn in red to say so. Chat opens the conversation list and, from there, "
            "the screen in Figure 5.9. Profile opens the screen in Figure 5.10. Those three "
            "tabs remain in place for the rest of the teacher session."
        )
        self.phone("teacher_home", "5.5",
                   "Teacher home, My Class Schedules. The plus button adds a class; the pencil edits; the bin deletes; the bottom bar switches tab.")

        self.heading3("5.3.4  Add Class Schedule")
        self.para(
            "The form that creates a class is titled ADD CLASS SCHEDULE and sits on a dark "
            "navy background. It is the most important data-entry page in the application, "
            "because every field on it is later used either to display a card or to decide "
            "which students may see that card."
        )
        self.para(
            "The teacher works from top to bottom. Select Department opens the list in "
            "Table 4.6 — Computer Science, Business Administration, Electrical Engineering, "
            "Mathematics, Physics, English. Select Section opens A, B, C or D. Subject Name "
            "is typed, for example Calculus. Select Semester opens 1 through 8. Select Day "
            "opens the seven weekdays. Select class date opens a calendar. Select class time "
            "opens a clock. At the bottom, Submit Schedule writes the row if the form is "
            "complete."
        )
        self.para(
            "If the teacher taps Submit Schedule before choosing a date or a time, the form "
            "does not close. A red sentence appears under the empty field: “Date is required” "
            "or “Time is required”. Figure 5.6 shows that protection in action. The teacher "
            "must finish the calendar and the clock, then submit again. On success the form "
            "disappears and the new card is waiting on My Class Schedules."
        )
        self.phone("add_form", "5.6",
                   "Complete Add Class Schedule form. Red sentences appear when date or time has been left empty.")
        self.para(
            "Figures 5.7 and 5.8 show the same form with its lists open, so that the reader "
            "can see the values rather than only the closed boxes. Department is a named "
            "college list, not the letters A to D; those letters belong to Section. Semester "
            "is a number of a degree programme, not a calendar year. Day is a weekday; date "
            "is a particular day of a month, chosen on a calendar for August 2026 in the "
            "photograph; time is a clock with an AM / PM switch."
        )
        self.phones_row(["add_dept", "add_section", "add_semester", "add_day"], "5.7",
                        "Four lists on Add Class Schedule: department, section (A–D), semester (1–8), and weekday.",
                        width=1.48)
        self.phones_row(["add_date", "add_time"], "5.8",
                        "Date picker and time picker used while adding a class. Cancel discards the choice; OK writes it into the form.",
                        width=2.15)
        self.para(
            "Once a class exists, the teacher does not stay on this form. The ordinary life "
            "of the class is on the home cards, in the chat, and on the profile."
        )

        self.heading3("5.3.5  Teacher chat")
        self.para(
            "From the Chat tab, or by opening a thread that a student has already started, "
            "the teacher reaches a conversation screen. Figure 5.9 shows a thread with a "
            "student named Awais Khan. The navy bar at the top holds a back arrow, a small "
            "portrait, the student’s name, and the word Offline. The back arrow returns to "
            "the list of conversations, not out of the application."
        )
        self.para(
            "Messages from the student arrive as pale bubbles on the left. Messages from "
            "the teacher leave as navy bubbles on the right. Each bubble carries a time. In "
            "the photograph the student has written a greeting; the teacher has answered "
            "with a greeting, a question after the student’s health, and an offer of help. "
            "At the bottom a box reads “Type a message…”. Beside it a red circular button "
            "with a paper-plane icon sends the line. Sending does not change screen; it only "
            "adds a bubble."
        )
        self.para(
            "This is the same conversation that the student sees from the other side in "
            "Figure 5.14. The words are identical; only the left and right sides are swapped, "
            "because each person sees his or her own lines on the right."
        )
        self.phone("teacher_chat", "5.9",
                   "Teacher chatting with the student Awais Khan. Incoming lines are grey and left; outgoing lines are navy and right.")

        self.heading3("5.3.6  Teacher profile")
        self.para(
            "The Profile tab opens a page titled Teacher Profile. A circular placeholder "
            "stands for a photograph. Under it the name Zubair and the email zubair@gmail.com "
            "identify the signed-in teacher. A white card then lists the professional facts "
            "that were typed at registration: Department Computer Science, Qualification PhD, "
            "Subjects Computer science, Gender Male, Phone 03039639571. Each row has a small "
            "icon so that the list can be scanned quickly."
        )
        self.para(
            "Beneath the card four actions are offered. Edit Profile opens a form in which "
            "those facts can be changed; after saving, the same profile page is shown with "
            "the new values. Terms and Conditions and About Us open reading pages about the "
            "application. Log Out, drawn in red, ends the session and returns the user to "
            "the role-selection screen. After log-out, My Class Schedules cannot be opened "
            "again without signing in."
        )
        self.phones_row(["teacher_profile", "teacher_profile_menu"], "5.10",
                        "Teacher profile. The left photograph shows the stored facts; the right photograph shows Edit Profile, legal pages, and Log Out.")

        self.heading2("5.4  Student flow")
        self.heading3("5.4.1  Student sign-in")
        self.para(
            "Tapping Student on the role-selection screen opens a page titled Student Sign In. "
            "Its shape is the same as the teacher’s sign-in, so that a person who has used one "
            "role is not surprised by the other. A back arrow returns to role selection. Two "
            "fields take email and password. Sign In, if the account is a registered student, "
            "opens My Classes (Figure 5.13). Sign In Using Google is offered as an alternative. "
            "“New here? Sign Up” opens the student registration form in Figure 5.12."
        )
        self.phone("student_signin", "5.11",
                   "Student sign-in. Success opens My Classes; Sign Up opens student registration.")

        self.heading3("5.4.2  Student sign-up")
        self.para(
            "The page is titled Student Sign Up, with the subtitle Student Registration. "
            "The form is shorter than the teacher’s form because a student is described by "
            "academic place rather than by qualification and subjects."
        )
        self.bullet("Full Name — for example Awais Khan.")
        self.bullet("Email address — the future sign-in address.")
        self.bullet("Select Department — the same college list used by teachers.")
        self.bullet("Select Field — programmes that belong to the chosen department, for example Software Engineering under Computer Science.")
        self.bullet("Select Section — A, B, C or D.")
        self.bullet("Select Semester — 1 to 8.")
        self.bullet("Password — hidden, with an eye icon.")
        self.para(
            "These four academic choices are not decoration. They are the keys that unlock "
            "the timetable. A student who registers in Computer Science, Software Engineering, "
            "section B, semester 8 will, after sign-in, see exactly those class cards that a "
            "teacher published with the same four values. Register, when the form is valid, "
            "creates the account and opens My Classes. “Already have an account? Sign In” "
            "returns to Figure 5.11."
        )
        self.phone("student_signup", "5.12",
                   "Student sign-up. Department, section and semester later decide which class cards appear on My Classes.")

        self.heading3("5.4.3  Student home — My Classes")
        self.para(
            "My Classes is the student’s home and the first tab of the bottom bar. The title "
            "sits at the top left; a refresh arrow sits at the top right. Each class is a "
            "maroon-to-navy card. Unlike the teacher’s card, the student’s card names the "
            "teacher and offers a speech-bubble instead of a pencil and a bin. A student "
            "cannot edit or delete a class. He can only read it and write to the teacher."
        )
        self.para(
            "In Figure 5.13 two cards are visible. The first is Calculus, taught by Zubair, "
            "on Monday, 24 August 2026, at 3:19 PM, for semester 8, Department Computer "
            "Science, Field Software Engineering, Section B. The second is Islamiyat, also "
            "taught by Zubair, on 25 August 2026 at 2:00 PM, with the same academic place. "
            "These are the same two meetings that the teacher published in Figure 5.5. The "
            "application has not made a second copy for the student; it has shown the same "
            "records through the matching rule drawn on the entity-relationship diagram."
        )
        self.para(
            "Tapping the speech-bubble on a card opens a chat with that named teacher "
            "(Figure 5.14). Tapping Chat on the bottom bar opens the list of all such "
            "threads. Tapping Profile opens Figure 5.15. Tapping the refresh arrow asks "
            "again for matching classes, which is how a student discovers a newly published "
            "or newly deleted meeting."
        )
        self.phone("student_home", "5.13",
                   "Student home, My Classes. Each card shows the teacher and a chat shortcut; there is no edit or delete control.")

        self.heading3("5.4.4  Student chat")
        self.para(
            "Figure 5.14 is the other face of Figure 5.9. The navy bar now shows the teacher’s "
            "name, Zubair, and the word Offline. The back arrow returns to the previous list. "
            "The student’s own lines — a greeting, a notice that he cannot attend tomorrow, "
            "and the reason that he is unwell in hospital — sit on the right in navy bubbles. "
            "The teacher’s replies sit on the left in grey. The student is using the chat for "
            "its intended academic purpose: to tell the teacher of a class about an absence, "
            "without asking for a personal phone number."
        )
        self.para(
            "The box “Type a message…” and the red send button work as they do on the "
            "teacher’s side. After sending, the student remains on this screen. The thread "
            "is stored, so that tomorrow both people will still see the same history."
        )
        self.phone("student_chat", "5.14",
                   "Student chatting with the teacher Zubair. The student is informing the teacher of an absence due to illness.")

        self.heading3("5.4.5  Student profile")
        self.para(
            "The Profile tab opens a page titled Student Profile. The signed-in person in "
            "the photograph is Awais Khan, email awaiskhan@gmail.com. A card lists the four "
            "academic facts that drive the timetable: Department Computer Science, Field "
            "Software Engineering, Section B, Semester 8. Those are the same four facts that "
            "appear at the foot of the class cards on My Classes. If the student were to edit "
            "the profile and move to another section, the home list would change on the next "
            "refresh, because the matching rule would point at a different set of schedules."
        )
        self.para(
            "Edit Profile, reached from the row under the card, is how those facts are "
            "changed. The bottom bar still offers Home and Chat. Signing out, as on the "
            "teacher side, returns the user to the role-selection screen."
        )
        self.phone("student_profile", "5.15",
                   "Student profile of Awais Khan. The four academic rows are the filter that builds My Classes.")

        self.heading2("5.5  A complete journey, told as a story")
        self.para(
            "It may help to retell the same screens as one day in the life of the two people "
            "who appear in the photographs. Zubair, a teacher in Computer Science, opens "
            "Islamian App. He sees the college emblem, then the role question, then he signs "
            "in as a teacher. On My Class Schedules he taps plus, chooses Computer Science, "
            "Software Engineering, section B, types Calculus, chooses semester 8, Monday, "
            "the 24th of August, and 3:19 PM. He submits. He repeats the form for Islamiyat "
            "on the 25th. Two cards now sit on his home."
        )
        self.para(
            "Awais Khan, a student of Software Engineering, section B, semester 8, opens the "
            "same application, chooses Student, and signs in. My Classes already shows "
            "Calculus and Islamiyat, both taught by Zubair. He taps the bubble on Calculus "
            "and writes that he is ill and cannot come. Zubair, on his own chat screen, "
            "reads the request and answers. Later each of them may open Profile to confirm "
            "that the stored facts are still correct, or Log Out when the phone is to be put away."
        )
        self.para(
            "No step in that story requires a file name, a programming word, or a database "
            "command. It is the behaviour of the product, and it is what the diagrams of "
            "Chapter 4 were drawn to support."
        )

        self.heading2("5.6  Summary")
        self.para(
            "The implemented system is a short path with a fork. Splash leads to role "
            "selection. Role selection leads to a sign-in. Sign-in leads to a three-tab "
            "home. For the teacher the home is a list he owns; for the student it is a list "
            "he is allowed to see. Chat joins the two lists. Profile describes the person "
            "who is signed in. The photographs in this chapter are that path, taken from a "
            "real device, in the same order a new user would walk it."
        )

    def ch6(self):
        self.page_break()
        self.heading1("SIX", "Testing")
        self.heading2("6.1  Purpose of testing")
        self.para(
            "Testing in this project was practical. Each screen of Chapter 5 was opened on "
            "a phone, used as a teacher or as a student, and judged against the requirements "
            "of Chapter 3. The question asked of every case was not whether a line of source "
            "was elegant, but whether the next screen was the one the prose had promised."
        )
        self.heading2("6.2  Test cases")
        self.table("6.1", "Principal test cases and their results",
                   ["ID", "Action taken", "Expected next screen or result", "Outcome"],
                   [
                       ["TC-01", "Launch the application.", "Splash, then role selection (when signed out).", "Passed"],
                       ["TC-02", "Tap Teacher, then Sign Up, fill the form, Register.", "My Class Schedules.", "Passed"],
                       ["TC-03", "Tap Student, then Sign Up with department, field, section, semester.", "My Classes.", "Passed"],
                       ["TC-04", "Sign in with a wrong password.", "Remain on sign-in; no home tab.", "Passed"],
                       ["TC-05", "Teacher submits a schedule with empty date.", "Red sentence “Date is required”; form stays open.", "Passed"],
                       ["TC-06", "Teacher submits a complete Calculus class for CS / SE / B / 8.", "Card appears on teacher home; matching student sees it.", "Passed"],
                       ["TC-07", "Teacher deletes a class.", "Card disappears; student list no longer shows it after refresh.", "Passed"],
                       ["TC-08", "Student taps the chat icon on a class card.", "Chat opens with that teacher’s name.", "Passed"],
                       ["TC-09", "Student sends a message; teacher opens the same thread.", "Words appear on the opposite side with a time.", "Passed"],
                       ["TC-10", "Teacher taps Log Out.", "Role-selection screen; home is not reachable without sign-in.", "Passed"],
                   ])
        self.heading2("6.3  Observations")
        self.para(
            "The matching rule is the most important behaviour to get right. A student "
            "registered in section A must not see a class published for section B, even when "
            "department and semester are the same. That check was tried by changing only the "
            "section on a test account and refreshing My Classes. The list emptied or filled "
            "as expected."
        )
        self.para(
            "Usability was judged by watching whether a new user reached for the plus button "
            "and the drop-down lists without instruction. The remaining difficulty is the "
            "number of fields on Add Class Schedule; the protection messages on date and time "
            "prevented incomplete rows during the trials."
        )
        self.heading2("6.4  Summary")
        self.para(
            "The ten cases cover the fork of roles, the two registrations, the unhappy paths "
            "of a bad password and an incomplete form, the happy path of a published class, "
            "the chat that follows it, and a clean log-out. They are sufficient to claim that "
            "the screens of Chapter 5 behave as described."
        )

    def ch7(self):
        self.page_break()
        self.heading1("SEVEN", "Conclusion and Future Work")
        self.heading2("7.1  Conclusion")
        self.para(
            "Islamian App was written to answer a small campus problem with a complete, "
            "visible product. A teacher can publish a class in a structured form. A student "
            "can see only the classes that belong to him or her. Either person can speak to "
            "the other inside the same application. The entity-relationship diagram, the "
            "data tables and the photographs all describe that same product. The writing of "
            "this thesis has been kept on the side of explanation rather than on the side of "
            "source listings, so that a reader who is not a programmer can still follow the "
            "journey from the college emblem to a sent message."
        )
        self.heading2("7.2  Future work")
        self.bullet("Push notifications so that a new class or a new message is noticed even when the application is closed.")
        self.bullet("A college administrator who can load a full semester timetable at once.")
        self.bullet("Attendance marking against a published class.")
        self.bullet("A room or building name on each class card.")
        self.bullet("A richer first-launch tutorial before the role-selection screen.")
        self.heading2("7.3  Closing remark")
        self.para(
            "If the application is used as intended, a student who is unwell should never "
            "again have to ask a classmate for a teacher’s personal number in order to miss "
            "a class with dignity. That is a modest aim, and it is the aim this project set "
            "out to meet."
        )

    def references(self):
        self.page_break()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("REFERENCES")
        set_run_font(r, size=16, bold=True, color=NAVY)
        refs = [
            "Pressman, R. S. and Maxim, B. R., Software Engineering: A Practitioner’s Approach, 8th ed., McGraw-Hill, 2014.",
            "Sommerville, I., Software Engineering, 10th ed., Pearson, 2015.",
            "Elmasri, R. and Navathe, S. B., Fundamentals of Database Systems, 7th ed., Pearson, 2016.",
            "Date, C. J., An Introduction to Database Systems, 8th ed., Addison-Wesley, 2003.",
            "Gamma, E., Helm, R., Johnson, R. and Vlissides, J., Design Patterns, Addison-Wesley, 1994.",
            "Fielding, R. T., Architectural Styles and the Design of Network-based Software Architectures, Ph.D. dissertation, University of California, Irvine, 2000.",
            "Google, Firebase Documentation: Authentication, Cloud Firestore and Cloud Storage, https://firebase.google.com/docs (accessed 2026).",
            "Google, Flutter Documentation, https://docs.flutter.dev (accessed 2026).",
            "Nielsen, J., Usability Engineering, Morgan Kaufmann, 1993.",
            "Preece, J., Rogers, Y. and Sharp, H., Interaction Design, 5th ed., Wiley, 2019.",
            "IEEE, IEEE Standard for Software and System Test Documentation, IEEE Std 829, 2008.",
            "Islamia College Peshawar, Department of Computer Science, undergraduate project guidelines, Peshawar.",
        ]
        for i, ref in enumerate(refs, 1):
            pp = self.doc.add_paragraph()
            pp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pp.paragraph_format.left_indent = Cm(1.25)
            pp.paragraph_format.first_line_indent = Cm(-1.25)
            pp.paragraph_format.space_after = Pt(8)
            pp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            run = pp.add_run(f"[{i}]  {ref}")
            set_run_font(run, size=12)

    def appendix(self):
        self.page_break()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("APPENDIX A")
        set_run_font(r, size=14, bold=True, color=NAVY)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("GLOSSARY OF SCREENS")
        set_run_font(r, size=16, bold=True, color=NAVY)
        self.para(
            "This glossary repeats the names printed on the photographs so that an examiner "
            "can find a screen by title.",
            first_indent=False,
        )
        rows = [
            ["Splash", "College emblem. No button. Next: role selection, or a home if already signed in."],
            ["YOU ARE TEACHER OR STUDENT", "Role fork. Teacher → teacher sign-in. Student → student sign-in."],
            ["Sign In (teacher)", "Email, password, Google, link to Teacher Sign Up."],
            ["Teacher Sign Up", "Professional form. Register → My Class Schedules."],
            ["My Class Schedules", "Teacher home. Plus, pencil, bin. Tabs: Home, Chat, Profile."],
            ["ADD CLASS SCHEDULE", "Form with lists, calendar, clock, Submit Schedule."],
            ["Chat (teacher)", "Thread with a named student. Send stays on the same screen."],
            ["Teacher Profile", "Facts, Edit Profile, Log Out."],
            ["Student Sign In", "Email, password, Google, link to Student Sign Up."],
            ["Student Sign Up", "Academic form. Register → My Classes."],
            ["My Classes", "Student home. Speech-bubble opens chat with the named teacher."],
            ["Chat (student)", "Thread with a named teacher."],
            ["Student Profile", "Academic facts that filter My Classes."],
        ]
        self.table("A.1", "Screen titles and what they lead to",
                   ["Title on the screen", "Meaning and next step"], rows)

        self.page_break()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("APPENDIX B")
        set_run_font(r, size=14, bold=True, color=NAVY)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("RELATIONSHIPS AMONG THE FIVE ENTITIES")
        set_run_font(r, size=16, bold=True, color=NAVY)
        self.table("B.1", "Relationships drawn on Figure 4.4",
                   ["From", "To", "Type", "Meaning in the application"],
                   [
                       ["Teacher", "Class Schedule", "One to many", "One teacher publishes many classes; each class has exactly one owner."],
                       ["Student", "Class Schedule", "Many to many (logical)", "A student sees every class whose department, field, section and semester equal the student’s own."],
                       ["Teacher", "Conversation", "One to many", "A teacher may be part of many student threads."],
                       ["Student", "Conversation", "One to many", "A student may be part of many teacher threads."],
                       ["Conversation", "Message", "One to many", "A thread contains an ordered list of lines."],
                   ])

    def build(self, out_path: Path):
        self.front_title_block()
        self.inner_title()
        self.certificate()
        self.dedication()
        self.acknowledgement()
        self.abstract()
        self.toc_placeholder()
        self.lists_of_tables_figures()
        # body — arabic page numbers
        self.new_section(header=True, fmt="decimal", start=1)
        self.ch1()
        self.ch2()
        self.ch3()
        self.ch4()
        self.ch5()
        self.ch6()
        self.ch7()
        self.references()
        self.appendix()
        self.doc.save(str(out_path))
        print(f"Wrote {out_path}")


def main():
    copy_screenshots()
    out = THESIS / "Awais_Khan_Thesis_Islamian_App.docx"
    Thesis().build(out)


if __name__ == "__main__":
    main()
